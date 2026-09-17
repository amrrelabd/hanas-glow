from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.management.base import BaseCommand

from inventory.models import Customer, Product, StockMovement, Supplier


def clean_str(value):
    if value is None:
        return ''
    return str(value).strip()


def clean_decimal(value):
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation:
        return None


def clean_int(value):
    if value is None or value == '':
        return 0
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return 0


def clean_phone(value):
    if value is None:
        return ''
    if isinstance(value, float):
        return str(int(value))
    return str(value).strip()


class Command(BaseCommand):
    help = 'Import Suppliers, Products, and Customers from the Hana\'s Glow BMS Excel file'

    def add_arguments(self, parser):
        parser.add_argument('file_path', type=str)

    def handle(self, *args, **options):
        wb = openpyxl.load_workbook(options['file_path'], data_only=True)

        self.import_suppliers(wb)
        self.import_products(wb)
        self.import_customers(wb)

    def import_suppliers(self, wb):
        ws = wb['\U0001f3ed Suppliers']
        created = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = clean_str(row[1])
            if not name:
                continue
            phone = clean_phone(row[2])
            _, was_created = Supplier.objects.get_or_create(name=name, defaults={'phone': phone})
            if was_created:
                created += 1
        self.stdout.write(self.style.SUCCESS(f'Suppliers imported: {created}'))

    def import_products(self, wb):
        ws = wb['\U0001f4e6 Products']
        created = 0
        skipped = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = clean_str(row[0])
            name = clean_str(row[2])
            if not sku or not name:
                continue
            if Product.objects.filter(sku=sku).exists():
                skipped += 1
                continue

            supplier = None
            supplier_name = clean_str(row[10])
            if supplier_name:
                supplier = Supplier.objects.filter(name__iexact=supplier_name).first()

            cost_price = clean_decimal(row[5]) or Decimal('0')
            sell_price = clean_decimal(row[6]) or Decimal('0')
            discount_price = clean_decimal(row[7])
            opening_stock = clean_int(row[8])

            product = Product.objects.create(
                name=name,
                sku=sku,
                category=clean_str(row[3]),
                brand=clean_str(row[1]),
                shade_color=clean_str(row[4]),
                supplier=supplier,
                cost_price=cost_price,
                sell_price=sell_price,
                discount_price=discount_price,
                min_stock_alert=clean_int(row[9]),
                barcode=clean_str(row[11]),
                notes=clean_str(row[12]),
            )

            if opening_stock:
                StockMovement.objects.create(
                    product=product,
                    quantity_delta=opening_stock,
                    reason=StockMovement.Reason.ADJUSTMENT,
                    reference_id='import',
                    notes='Opening stock from Excel import',
                )
            created += 1
        self.stdout.write(self.style.SUCCESS(f'Products imported: {created} (skipped {skipped} existing)'))

    def import_customers(self, wb):
        ws = wb['\U0001f465 Customers']
        created = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = clean_str(row[1])
            if not name:
                continue
            phone = clean_phone(row[2])
            _, was_created = Customer.objects.get_or_create(
                name=name,
                defaults={
                    'phone': phone,
                    'city': clean_str(row[3]),
                    'address': clean_str(row[4]),
                    'notes': clean_str(row[6]),
                },
            )
            if was_created:
                created += 1
        self.stdout.write(self.style.SUCCESS(f'Customers imported: {created}'))
